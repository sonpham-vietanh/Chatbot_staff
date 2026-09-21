import io

import pytest
from docx import Document

from app.services.knowledge_ingest import KnowledgeIngestService


class FakeAdminService:
    def __init__(self):
        self.calls = []

    def create_note(self, title, department, content, status, created_by):
        note = {"id": "fake-id", "title": title, "department": department,
                "content": content, "status": status, "created_by": created_by}
        self.calls.append(note)
        return note


def test_ingest_text_creates_approved_note():
    admin = FakeAdminService()
    result = KnowledgeIngestService(admin).ingest(
        "quy_trinh_moi.txt",
        "Bước một: gửi đề nghị cho HR.".encode("utf-8"),
        "HR",
    )

    assert result["status"] == "approved"
    assert admin.calls[0]["status"] == "approved"
    assert admin.calls[0]["department"] == "HR"
    assert "Bước một" in admin.calls[0]["content"]


def test_ingest_docx_converts_word_headings_to_markdown():
    document = Document()
    document.add_paragraph("Nội quy công ty", style="Title")
    document.add_paragraph("Chương 1: Giờ làm việc", style="Heading 1")
    document.add_paragraph("Nhân viên làm việc 8 tiếng mỗi ngày.")
    document.add_paragraph("Mục 1.1: Ca sáng", style="Heading 2")
    document.add_paragraph("Bắt đầu từ 8h00.")
    buffer = io.BytesIO()
    document.save(buffer)

    admin = FakeAdminService()
    result = KnowledgeIngestService(admin).ingest("noi_quy.docx", buffer.getvalue(), "HR")

    content = admin.calls[0]["content"]
    assert "# Nội quy công ty" in content
    assert "## Chương 1: Giờ làm việc" in content
    assert "### Mục 1.1: Ca sáng" in content
    assert result["status"] == "approved"


def test_ingest_docx_extracts_table_content():
    document = Document()
    document.add_paragraph("Chính sách học bổng", style="Heading 1")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Mức học bổng"
    table.rows[0].cells[1].text = "Điều kiện"
    table.rows[1].cells[0].text = "50%"
    table.rows[1].cells[1].text = "GPA >= 3.5"
    document.add_paragraph("Áp dụng từ năm 2026.")
    buffer = io.BytesIO()
    document.save(buffer)

    admin = FakeAdminService()
    KnowledgeIngestService(admin).ingest("hoc_bong.docx", buffer.getvalue(), "HR")

    content = admin.calls[0]["content"]
    assert "# Chính sách học bổng" in content
    assert "Mức học bổng" in content and "Điều kiện" in content
    assert "50%" in content and "GPA >= 3.5" in content
    assert "Áp dụng từ năm 2026." in content


class FakeLLM:
    def describe_image(self, image_bytes, mime_type):
        return "Bảng lương: bậc 1 = 5 triệu, bậc 2 = 7 triệu."

    def answer(self, question, contexts, history=None):
        raise NotImplementedError


def test_ingest_docx_describes_images_via_llm():
    document = Document()
    document.add_paragraph("Bảng lương nhân viên", style="Heading 1")
    document.add_picture(io.BytesIO(_tiny_png()))
    buffer = io.BytesIO()
    document.save(buffer)

    admin = FakeAdminService()
    KnowledgeIngestService(admin, FakeLLM()).ingest("bang_luong.docx", buffer.getvalue(), "HR")

    content = admin.calls[0]["content"]
    assert "Nội dung trích xuất từ hình ảnh" in content
    assert "bậc 1 = 5 triệu" in content


def _tiny_png() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0"
        b"\x00\x00\x03\x01\x01\x00\x18\xdd\x8d\xb0\x00\x00\x00\x00IEND\xaeB`\x82"
    )


def test_ingest_rejects_unsupported_file():
    with pytest.raises(ValueError, match="Định dạng chưa được hỗ trợ"):
        KnowledgeIngestService(FakeAdminService()).ingest("secret.exe", b"data", "HR")


def test_ingest_rejects_oversized_file():
    admin = FakeAdminService()
    oversized = b"x" * (10 * 1024 * 1024 + 1)
    with pytest.raises(ValueError, match="10 MB"):
        KnowledgeIngestService(admin).ingest("big.txt", oversized, "HR")
