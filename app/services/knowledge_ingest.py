from __future__ import annotations

import io
import json
from pathlib import Path
import re
from typing import Any

import frontmatter

from app.services.admin_service import AdminService
from app.services.llm import LLMProvider

ALLOWED_EXTENSIONS = {".md", ".txt", ".csv", ".json", ".pdf", ".docx"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMAGES_PER_FILE = 15
"""Chặn trần số ảnh gửi qua AI Vision mỗi lần upload — tránh 1 file docx nhiều ảnh
trang trí làm request treo lâu hoặc phát sinh chi phí AI vượt kiểm soát."""


class KnowledgeIngestService:
    """Trích nội dung văn bản (+ mô tả ảnh qua AI Vision nếu có LLM) từ file HR/admin
    upload, tạo thành note trong Supabase."""

    def __init__(self, admin_service: AdminService, llm: LLMProvider | None = None):
        self.admin_service = admin_service
        self.llm = llm

    def ingest(self, filename: str, content: bytes, department: str | None, title: str | None = None) -> dict[str, Any]:
        safe_name = self._safe_filename(filename)
        extension = Path(safe_name).suffix.casefold()
        if extension not in ALLOWED_EXTENSIONS:
            raise ValueError(f"Định dạng chưa được hỗ trợ: {extension or 'không có phần mở rộng'}")
        if len(content) > MAX_UPLOAD_BYTES:
            raise ValueError("File vượt quá giới hạn 10 MB")

        note_title = title.strip() if title and title.strip() else Path(safe_name).stem
        extracted_text = self._extract_text(safe_name, content).strip()
        body = extracted_text or f"File gốc: `{safe_name}`. Chờ HR bổ sung nội dung có thể tìm kiếm."

        duplicate = self._find_possible_duplicate(note_title)

        note = self.admin_service.create_note(
            title=note_title,
            department=department or "Unassigned",
            content=body,
            status="approved",
            created_by="HR_Upload",
        )
        result = {
            "note_id": note["id"],
            "status": "approved",
            "title": note["title"],
            "source_file": safe_name,
            "message": "Đã tạo note và duyệt tự động — chatbot có thể trả lời từ nội dung này ngay.",
        }
        if duplicate:
            result["warning"] = (
                f"Tên gần giống note đã có: \"{duplicate['title']}\" (đang {duplicate['status']}). "
                "Kiểm tra lại tránh dữ liệu trùng/mâu thuẫn — nếu đây là bản cập nhật, nên từ chối "
                "hoặc xoá note cũ."
            )
        return result

    def _find_possible_duplicate(self, title: str) -> dict[str, Any] | None:
        """So khớp tên đơn giản (không phân biệt hoa/thường, chứa lẫn nhau) với các note đã
        duyệt — không có embedding similarity ở đây vì chỉ cần cảnh báo nhanh trước khi ghi,
        không cần chính xác tuyệt đối."""
        normalized = title.strip().casefold()
        if not normalized:
            return None
        for note in self.admin_service.list_notes("approved"):
            existing = (note.get("title") or "").strip().casefold()
            if existing and (existing == normalized or existing in normalized or normalized in existing):
                return note
        return None

    @staticmethod
    def _safe_filename(filename: str) -> str:
        name = Path(filename or "upload").name
        name = re.sub(r"[^\w.() -]+", "_", name, flags=re.UNICODE).strip(" .")
        return name[:160] or "upload"

    def _extract_text(self, filename: str, content: bytes) -> str:
        extension = Path(filename).suffix.casefold()
        if extension in {".md", ".txt", ".csv", ".json"}:
            text = content.decode("utf-8-sig", errors="replace")
            if extension == ".md":
                try:
                    parsed = frontmatter.loads(text)
                    return parsed.content
                except Exception:
                    return text
            if extension == ".json":
                try:
                    return json.dumps(json.loads(text), ensure_ascii=False, indent=2)
                except json.JSONDecodeError:
                    return text
            return text
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if extension == ".docx":
            from docx import Document
            from docx.oxml.table import CT_Tbl
            from docx.oxml.text.paragraph import CT_P
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            document = Document(io.BytesIO(content))
            lines = []
            for child in document.element.body.iterchildren():
                if isinstance(child, CT_P):
                    paragraph = Paragraph(child, document)
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    level = KnowledgeIngestService._docx_heading_level(paragraph.style.name if paragraph.style else "")
                    lines.append(f"{'#' * level} {text}" if level else text)
                elif isinstance(child, CT_Tbl):
                    rendered = KnowledgeIngestService._render_docx_table(Table(child, document))
                    if rendered:
                        lines.append(rendered)
            image_section = self._describe_docx_images(document)
            if image_section:
                lines.append(image_section)
            return "\n\n".join(lines)
        return ""

    def _describe_docx_images(self, document) -> str:
        """Docx là file zip nên ảnh nằm tách rời trong package, không đọc được qua
        document.paragraphs — phải lấy riêng từ các relationship kiểu 'image' rồi gửi qua
        AI Vision để không bỏ sót nội dung (bảng số liệu chụp ảnh, sơ đồ có chữ...)."""
        if not self.llm:
            return ""
        images = self._extract_docx_images(document)
        if not images:
            return ""
        described = []
        for index, (blob, content_type) in enumerate(images[:MAX_IMAGES_PER_FILE], start=1):
            try:
                description = (self.llm.describe_image(blob, content_type) or "").strip()
            except Exception:
                description = ""
            if description:
                described.append(f"**Hình {index}:** {description}")
        if not described:
            return ""
        section = "## Nội dung trích xuất từ hình ảnh (AI mô tả tự động)\n\n" + "\n\n".join(described)
        if len(images) > MAX_IMAGES_PER_FILE:
            section += f"\n\n_(Chỉ xử lý {MAX_IMAGES_PER_FILE}/{len(images)} ảnh đầu tiên trong file.)_"
        return section

    @staticmethod
    def _extract_docx_images(document) -> list[tuple[bytes, str]]:
        images = []
        seen_parts = set()
        for rel in document.part.rels.values():
            if rel.reltype.endswith("/image") and rel.target_part.partname not in seen_parts:
                seen_parts.add(rel.target_part.partname)
                images.append((rel.target_part.blob, rel.target_part.content_type))
        return images

    @staticmethod
    def _render_docx_table(table) -> str:
        """python-docx không đọc bảng qua document.paragraphs — phải duyệt riêng, nếu không
        nội dung trong bảng (vd bảng % học bổng, bậc lương) sẽ biến mất khi trích xuất."""
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            return ""
        header, *body = rows
        lines = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
        lines.extend("| " + " | ".join(row) + " |" for row in body)
        return "\n".join(lines)

    @staticmethod
    def _docx_heading_level(style_name: str) -> int:
        """Word lưu cấu trúc mục/chương qua style ('Heading 1'..'Heading 9', 'Title') —
        chuyển thành Markdown '#'..'####' để chunk_content() cắt đúng theo mục mà không
        cần sửa tay lại nội dung sau khi upload."""
        name = (style_name or "").casefold()
        if name == "title":
            return 1
        match = re.match(r"heading (\d+)", name)
        if match:
            return min(int(match.group(1)) + 1, 4)
        return 0
